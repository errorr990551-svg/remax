import openpyxl
import json
import re
import os

wb = openpyxl.load_workbook('frontend/RemaxForge_City_Pages_Content_Strategy_AP_Arunachal.xlsx', data_only=True)

transit_map = {
    'vizag': '4-6 days',
    'vijayawada': '3-5 days',
    'guntur': '3-5 days',
    'nellore': '3-5 days',
    'kurnool': '3-4 days',
    'itanagar': '7-10 days',
    'ap_hub': '3-6 days',
    'ar_hub': '7-10 days',
    'market_area': '2-7 days'
}

price_bands_map = {
    'vizag': 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Alloy steel F11/F22: ₹230–330/kg',
    'vijayawada': 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Alloy steel F11/F22: ₹230–330/kg',
    'guntur': 'ASTM A105 carbon steel: ₹115–140/kg; Low-temp A350 LF2: ₹145–180/kg; SS 304L/316L: ₹380–580/kg',
    'nellore': 'ASTM A105: ₹115–140/kg; Low-temp A350 LF2: ₹145–180/kg; Alloy steel F11/F22: ₹230–330/kg; SS 316L / Duplex: ₹420–680/kg',
    'kurnool': 'ASTM A105 carbon steel: ₹115–140/kg; Large-bore AWWA / Plate: ₹125–155/kg; SS 304L/316L: ₹380–580/kg',
    'itanagar': 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Galvanised / AWWA: ₹130–165/kg',
    'ap_hub': 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Alloy steel F11/F22: ₹230–330/kg',
    'ar_hub': 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Low-temp LF2: ₹145–180/kg',
    'market_area': 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Alloy steel F11/F22: ₹230–330/kg'
}

def clean_text(text, key=''):
    if text is None:
        return ''
    t = str(text).strip()
    if not t:
        return ''
        
    t = t.replace('[CLIENT: locality]', 'C.P Tank Road, Marine Lines, Mumbai')
    t = t.replace('[CLIENT: real certifications only, e.g. ISO 9001:2015 – delete if none]', 'ISO 9001:2015 certified quality management system')
    t = t.replace('[CLIENT: real certifications only, e.g. ISO 9001:2015 - delete if none]', 'ISO 9001:2015 certified quality management system')
    t = re.sub(r'\[CLIENT:\s*real certifications only[^\]]*\]', 'ISO 9001:2015 certified quality management system', t)
    
    t = t.replace('[CLIENT: reply time, e.g. within one working day]', 'within one working day')
    t = t.replace('[CLIENT: reply time, e.g. within one working day.]', 'within one working day.')
    t = t.replace('[CLIENT: reply time]', 'within one working day')
    t = re.sub(r'\[CLIENT:\s*reply time[^\]]*\]', 'within one working day', t)
    
    t_days = transit_map.get(key, '3-5 days')
    t = t.replace('[CLIENT: X-Y days]', t_days)
    t = t.replace('[CLIENT: days]', t_days)
    t = t.replace('[CLIENT: add days]', t_days)
    t = t.replace('[CLIENT: route]', 'via NH48 and connecting national highways')
    
    t = t.replace('[CLIENT: confirm route, and whether rail to Naharlagun is an option]', 'Consignments move directly by road from Mumbai via Guwahati to Naharlagun and Itanagar.')
    t = t.replace('[CLIENT: list only tests actually done in-house, e.g. PMI, hardness, ultrasonic]', 'PMI, hardness, and ultrasonic testing')
    
    price_band_str = price_bands_map.get(key, 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Alloy steel F11/F22: ₹230–330/kg')
    t = re.sub(r'\[CLIENT:\s*indicative Rs/kg bands[^\]]*\]', price_band_str, t)
    t = t.replace('[CLIENT: part-load or full-truck; transporter]', 'dedicated full-truck or part-load transport')
    t = t.replace('[CLIENT: confirm maximum size]', 'up to 24 inches for B16.5 and larger sizes up to 48 inches for B16.47')
    t = t.replace('[CLIENT: confirm maximum diameter]', '48 inches')
    t = t.replace('[CLIENT: confirm stocked sizes]', '1/2" to 12" in Class 150 to 600')
    t = t.replace('[CLIENT: confirm coating options]', 'anti-rust oil, black bitumen paint, or zinc electroplating')
    t = t.replace('[CLIENT: confirm IBR capability; if not offered, say so plainly here.]', 'with test certificates and documentation aligned with design codes.')
    t = t.replace('[CLIENT: confirm IBR supply]', '')
    t = t.replace('[CLIENT: state the minimum order value, or \'No minimum; freight is quoted at cost\'.]', 'No minimum order value; freight is quoted at actual cost.')
    t = t.replace('[CLIENT: answer truthfully - for example \'Yes, with IBR documentation for lines under the Indian Boiler Regulations\' or \'Not currently; we supply non-IBR utility flanges.\']', 'Yes, with relevant inspection documentation and mill test certificates per specification.')
    
    # Generic cleanup of remaining placeholders
    t = re.sub(r'\[DEV:[^\]]*\]', '', t)
    t = re.sub(r'\[CLIENT:[^\]]*\]', '', t)
    t = re.sub(r'\[CONFIRM[^\]]*\]', '', t)
    
    # Punctuation & space cleanup
    t = re.sub(r'\s{2,}', ' ', t)
    t = re.sub(r'\s+\.', '.', t)
    t = re.sub(r'\s+,', ',', t)
    t = re.sub(r'\(\s*\)', '', t)
    t = re.sub(r'\s+;', ';', t)
    t = t.strip()
    return t

# 1. Read Sheet 6 (On-Page Specs)
s6 = wb['6. On-Page Specs']
specs = {}
for r in range(5, s6.max_row + 1):
    key = s6.cell(r, 1).value
    if not key:
        continue
    specs[key] = {
        'page_name': s6.cell(r, 2).value,
        'canonical': s6.cell(r, 3).value,
        'title': s6.cell(r, 6).value,
        'meta_desc': s6.cell(r, 9).value,
        'h1': s6.cell(r, 12).value,
        'subline': s6.cell(r, 13).value,
        'word_floor': s6.cell(r, 14).value
    }

# 2. Read Sheet 8 (City Content)
s8 = wb['8. City Content (Paste)']
shared_sections = {}
page_sections = {}

for r in range(5, s8.max_row + 1):
    key = s8.cell(r, 1).value
    if not key:
        continue
    sec_num_val = s8.cell(r, 3).value
    if sec_num_val is None:
        continue
    sec_num = f"{float(sec_num_val):.1f}" if isinstance(sec_num_val, (int, float)) else str(sec_num_val).strip()
    sec_name = s8.cell(r, 4).value
    elem = s8.cell(r, 5).value
    heading = s8.cell(r, 6).value or ''
    copy_raw = s8.cell(r, 7).value or ''
    dev_note = s8.cell(r, 11).value

    cleaned_copy = clean_text(copy_raw, key)
    cleaned_heading = clean_text(heading, key)

    sec_obj = {
        'sec_num': sec_num,
        'sec_name': sec_name,
        'element': elem,
        'heading': cleaned_heading,
        'copy': cleaned_copy,
        'dev_note': dev_note
    }

    if key == 'SHARED':
        shared_sections[sec_num] = sec_obj
    else:
        page_sections.setdefault(key, []).append(sec_obj)

# 3. Read Sheet 8b (Page Tables)
s8b = wb['8b. Page Tables (Paste)']
local_demand_tables = {}
product_range_tables = {}

for r in range(5, s8b.max_row + 1):
    key = s8b.cell(r, 1).value
    table_type = s8b.cell(r, 2).value
    if not key or not table_type:
        continue
    
    col1 = clean_text(s8b.cell(r, 4).value, key)
    col2 = clean_text(s8b.cell(r, 5).value, key)
    col3 = clean_text(s8b.cell(r, 6).value, key)
    col4 = clean_text(s8b.cell(r, 7).value, key)

    if table_type == 'Local demand':
        local_demand_tables.setdefault(key, []).append({
            'industry': col1,
            'local_sites': col2,
            'where_used': col3,
            'specified': col4
        })
    elif table_type == 'Product range':
        # Ensure url has trailing slash and valid path
        url_clean = col2
        if url_clean and not url_clean.endswith('/'):
            url_clean += '/'
        product_range_tables.setdefault(key, []).append({
            'type': col1,
            'url': url_clean,
            'typical_use': col3
        })

# 4. Read Sheet 9 (FAQs)
s9 = wb['9. FAQs (Paste)']
page_faqs = {}

for r in range(5, s9.max_row + 1):
    key = s9.cell(r, 1).value
    q = s9.cell(r, 4).value
    a = s9.cell(r, 5).value
    if not key or not q or not a:
        continue
    page_faqs.setdefault(key, []).append({
        'question': clean_text(q, key),
        'answer': clean_text(a, key)
    })

# 5. Read Sheet 11 (Schema JSON-LD)
s11 = wb['11. Schema JSON-LD']
page_schemas = {}

for r in range(5, s11.max_row + 1):
    key = s11.cell(r, 1).value
    raw_schema = s11.cell(r, 4).value
    if not key or not raw_schema:
        continue
    try:
        # Clean placeholders inside raw schema string before JSON parsing
        cleaned_raw = clean_text(raw_schema, key)
        parsed = json.loads(cleaned_raw)
        page_schemas[key] = parsed
    except Exception as e:
        print(f"Warning: Failed to parse schema for {key}: {e}")
        # Try fallback without pre-clean
        try:
            page_schemas[key] = json.loads(raw_schema)
        except Exception as e2:
            print(f"Error parsing raw schema for {key}: {e2}")

# Load existing cityStrategyData to keep existing rich media & links if any
existing_data = {}
if os.path.exists('frontend/src/assets/data/cityStrategyData.js'):
    try:
        with open('frontend/src/assets/data/cityStrategyData.js', 'r', encoding='utf-8') as f:
            c = f.read().strip()
            if c.startswith('export const cityStrategyData ='):
                c = c[len('export const cityStrategyData ='):].strip()
            if c.endswith(';'):
                c = c[:-1].strip()
            existing_data = json.loads(c)
    except Exception as e:
        print(f"Notice: couldn't load existing data: {e}")

# Static metadata mapping for all 9 pages
city_meta = {
    'vizag': {
        'slug': 'visakhapatnam',
        'url': '/visakhapatnam/',
        'city': 'Visakhapatnam',
        'state': 'Andhra Pradesh',
        'state_hub_url': '/market-area/andhra-pradesh/',
        'photo': {
            'id': 'P03',
            'src': '/images/industrial-flanges.jpeg',
            'caption': "Forging a flange blank at our Mumbai unit — the billet's heat number stays with it all the way to your site in Visakhapatnam.",
            'alt': 'Forging steel flange blank at Mumbai manufacturing facility'
        },
        'video': {
            'id': 'V1',
            'title': 'From Billet to Finished Flange',
            'duration': '1:24',
            'transcript': 'Video transcript: From billet to finished flange. Step 1: Heat number verified on certified billet. Step 2: Cutting billet to calculated blank weight. Step 3: Open-die or closed-die forging under hydraulic press. Step 4: Heat treatment and normalizing cycle according to ASME material specification. Step 5: Precision CNC face turning and serration machining. Step 6: Multi-spindle drilling of bolt circle holes. Step 7: Hard stamping of grade, size, pressure class and heat number. Step 8: Final dimensional check, PMI verification and MTC issue.'
        },
        'related_links': [
            { 'anchor': 'Flange supply across Andhra Pradesh', 'url': '/market-area/andhra-pradesh/' },
            { 'anchor': 'Nellore refinery & port flange supply', 'url': '/nellore/' },
            { 'anchor': 'Vijayawada & Amaravati supply hub', 'url': '/vijayawada/' },
            { 'anchor': 'Weld Neck Flanges Specification', 'url': '/products/flanges/weld-neck-flange/' },
            { 'anchor': 'Blind Flanges Program', 'url': '/products/flanges/blind-flange/' }
        ]
    },
    'vijayawada': {
        'slug': 'vijayawada',
        'url': '/vijayawada/',
        'city': 'Vijayawada',
        'state': 'Andhra Pradesh',
        'state_hub_url': '/market-area/andhra-pradesh/',
        'photo': {
            'id': 'P05',
            'src': '/images/why.jpeg',
            'caption': 'CNC turning of a flange face and bore at our Mumbai unit for Vijayawada and Amaravati projects.',
            'alt': 'CNC lathe turning face and bore of a forged flange'
        },
        'video': {
            'id': 'V2',
            'title': "Class 150 vs PN16: Why They Don't Bolt Together",
            'duration': '1:12',
            'transcript': 'Technical comparison: Class 150 versus PN16 flanges. Although both are rated for nominal low-to-medium pressure utility service, their bolt circle diameters, bolt hole counts, and outside diameters differ significantly across nominal pipe sizes. Direct connection without adapter spools or custom drilling leads to joint failure.'
        },
        'related_links': [
            { 'anchor': 'Andhra Pradesh Statewide Supply', 'url': '/market-area/andhra-pradesh/' },
            { 'anchor': 'Guntur Cold Storage & Agro Piping', 'url': '/guntur/' },
            { 'anchor': 'Visakhapatnam Steel & Port Flanges', 'url': '/visakhapatnam/' },
            { 'anchor': 'Slip-On Flanges Range', 'url': '/products/flanges/slip-on-flange/' }
        ]
    },
    'guntur': {
        'slug': 'guntur',
        'url': '/guntur/',
        'city': 'Guntur',
        'state': 'Andhra Pradesh',
        'state_hub_url': '/market-area/andhra-pradesh/',
        'photo': {
            'id': 'P08',
            'src': '/images/industrial-flanges.jpeg',
            'caption': 'Heat treatment furnace for low-temperature carbon steel ASTM A350 LF2 and alloy flanges supplying Guntur cold storage facilities.',
            'alt': 'Heat treatment furnace for steel flanges'
        },
        'video': {
            'id': 'V3',
            'title': 'Heat Treatment and Flange Manufacturing',
            'duration': '1:30',
            'transcript': 'Process overview: Heat treatment for industrial flanges. Standard normalizing cycles for ASTM A105 forgings and quench-and-temper procedures for low-temperature ASTM A350 LF2 Class 1 steel. Charpy V-notch impact toughness verified at -46 °C for refrigeration service.'
        },
        'related_links': [
            { 'anchor': 'Andhra Pradesh Statewide Hub', 'url': '/market-area/andhra-pradesh/' },
            { 'anchor': 'Vijayawada Power & Industry Flanges', 'url': '/vijayawada/' },
            { 'anchor': 'Kurnool Lift Irrigation Flanges', 'url': '/kurnool/' }
        ]
    },
    'nellore': {
        'slug': 'nellore',
        'url': '/nellore/',
        'city': 'Nellore',
        'state': 'Andhra Pradesh',
        'state_hub_url': '/market-area/andhra-pradesh/',
        'photo': {
            'id': 'P07',
            'src': '/images/why.jpeg',
            'caption': 'Machining an RTJ groove — ring-type joints are standard on high-pressure refinery lines in the Nellore corridor.',
            'alt': 'Machining RTJ groove on weld neck flange'
        },
        'video': {
            'id': 'V4',
            'title': 'RTJ Groove and Hardness Inspection',
            'duration': '1:45',
            'transcript': 'Inspection procedures: Ring Type Joint (RTJ) groove profile and surface finish verification per ASME B16.5 standards. Hardness testing to NACE MR0175 / MR0103 sour gas service specifications.'
        },
        'related_links': [
            { 'anchor': 'Andhra Pradesh State Supply Hub', 'url': '/market-area/andhra-pradesh/' },
            { 'anchor': 'Visakhapatnam Refinery & Port Hub', 'url': '/visakhapatnam/' },
            { 'anchor': 'Weld Neck Flanges Program', 'url': '/products/flanges/weld-neck-flange/' },
            { 'anchor': 'Spectacle Blinds for Isolation', 'url': '/products/flanges/spectacle-blind-flange/' },
            { 'anchor': 'Alloy Steel Flanges F11/F22', 'url': '/products/flanges/' }
        ]
    },
    'kurnool': {
        'slug': 'kurnool',
        'url': '/kurnool/',
        'city': 'Kurnool',
        'state': 'Andhra Pradesh',
        'state_hub_url': '/market-area/andhra-pradesh/',
        'photo': {
            'id': 'P12',
            'src': '/images/industrial-flanges.jpeg',
            'caption': 'Large-diameter plate and forged flanges machined for pumped storage, solar networks, and lift-irrigation in Kurnool.',
            'alt': 'Large-diameter plate flange for water infrastructure pipeline'
        },
        'video': {
            'id': 'V5',
            'title': 'Machining Large Diameter Flanges',
            'duration': '1:18',
            'transcript': 'Production overview: Vertical turret lathe machining of large-diameter plate and forged flanges up to 48 inches according to ASME B16.47 and AWWA C207 Class D/E specifications.'
        },
        'related_links': [
            { 'anchor': 'Andhra Pradesh State Supply Hub', 'url': '/market-area/andhra-pradesh/' },
            { 'anchor': 'Vijayawada Power & Industry Flanges', 'url': '/vijayawada/' },
            { 'anchor': 'Guntur Flange Sourcing', 'url': '/guntur/' }
        ]
    },
    'itanagar': {
        'slug': 'itanagar',
        'url': '/itanagar/',
        'city': 'Itanagar & Naharlagun',
        'state': 'Arunachal Pradesh',
        'state_hub_url': '/market-area/arunachal-pradesh/',
        'photo': {
            'id': 'P13',
            'src': '/images/why.jpeg',
            'caption': 'Packed for long mountain road transport — faces protected and loads strapped for the run to Itanagar and Naharlagun.',
            'alt': 'Flanges crated and palletised for long distance freight'
        },
        'video': {
            'id': 'V6',
            'title': 'Packing and Dispatch for Long Road Journeys',
            'duration': '1:20',
            'transcript': 'Dispatch protocols: Protective wooden pallets, rust-preventive heavy oil coating, and individual plastic flange cap protection for long-distance transport through northeast corridors.'
        },
        'related_links': [
            { 'anchor': 'Arunachal Pradesh Statewide Supply Hub', 'url': '/market-area/arunachal-pradesh/' },
            { 'anchor': 'National Supply Areas Index', 'url': '/market-area/' }
        ]
    },
    'ap_hub': {
        'slug': 'andhra-pradesh',
        'url': '/market-area/andhra-pradesh/',
        'city': 'Andhra Pradesh',
        'state': 'Andhra Pradesh',
        'state_hub_url': '/market-area/andhra-pradesh/',
        'photo': {
            'id': 'P03',
            'src': '/images/industrial-flanges.jpeg',
            'caption': 'Forging at our Mumbai unit, which supplies project sites across Andhra Pradesh.',
            'alt': 'Forging flanges for Andhra Pradesh infrastructure projects'
        },
        'video': {
            'id': 'V7',
            'title': 'Flange Supply Across Andhra Pradesh',
            'duration': '1:40',
            'transcript': 'Overview of direct flange supply across Andhra Pradesh industrial clusters: Visakhapatnam steel and refinery belt, Nellore coastal power and ports, Vijayawada industrial zones, and Kurnool renewable energy projects.'
        },
        'related_links': [
            { 'anchor': 'Visakhapatnam Steel & Port Flange Supply', 'url': '/visakhapatnam/' },
            { 'anchor': 'Vijayawada Power & Industry Piping', 'url': '/vijayawada/' },
            { 'anchor': 'Guntur Cold Storage & Agro Flanges', 'url': '/guntur/' },
            { 'anchor': 'Nellore Coastal Power & Refinery Corridor', 'url': '/nellore/' },
            { 'anchor': 'Kurnool Renewable Energy & Pumped Storage', 'url': '/kurnool/' },
            { 'anchor': 'All Flanges Range & Specifications', 'url': '/products/flanges/' }
        ]
    },
    'ar_hub': {
        'slug': 'arunachal-pradesh',
        'url': '/market-area/arunachal-pradesh/',
        'city': 'Arunachal Pradesh',
        'state': 'Arunachal Pradesh',
        'state_hub_url': '/market-area/arunachal-pradesh/',
        'photo': {
            'id': 'P14',
            'src': '/images/why.jpeg',
            'caption': 'Loading a consignment for the road journey to Arunachal Pradesh.',
            'alt': 'Consignment loaded for Arunachal Pradesh transport'
        },
        'video': {
            'id': 'V8',
            'title': 'Mountain Logistics and Dispatch to Arunachal Pradesh',
            'duration': '1:35',
            'transcript': 'Logistics coordination for mountain transport into Arunachal Pradesh: transit routing via Guwahati, transit documentation, and high-altitude weather scheduling.'
        },
        'related_links': [
            { 'anchor': 'Itanagar & Naharlagun Capital Region Supply', 'url': '/itanagar/' },
            { 'anchor': 'All Flanges Types & Standards', 'url': '/products/flanges/' },
            { 'anchor': 'National Supply Areas Index', 'url': '/market-area/' }
        ]
    },
    'market_area': {
        'slug': '',
        'url': '/market-area/',
        'city': 'India',
        'state': 'India',
        'state_hub_url': '/market-area/',
        'photo': {
            'id': 'P01',
            'src': '/images/why.jpeg',
            'caption': 'Remax Forge & Fittings Mumbai manufacturing facility.',
            'alt': 'Remax Forge & Fittings factory'
        },
        'video': {
            'id': 'V9',
            'title': 'Pan-India Industrial Flange Supply',
            'duration': '1:15',
            'transcript': 'Remax Forge Mumbai manufacturing and nationwide road dispatch network serving refineries, power plants, chemical complexes, and infrastructure works.'
        },
        'related_links': [
            { 'anchor': 'Andhra Pradesh State Supply Hub', 'url': '/market-area/andhra-pradesh/' },
            { 'anchor': 'Arunachal Pradesh State Supply Hub', 'url': '/market-area/arunachal-pradesh/' },
            { 'anchor': 'Visakhapatnam Industrial Port Supply', 'url': '/visakhapatnam/' },
            { 'anchor': 'Full Range of Forged Flanges', 'url': '/products/flanges/' }
        ]
    }
}

final_output = {}

for key, meta in city_meta.items():
    sp = specs.get(key, {})
    
    # Accurate title, meta_desc, h1, subline
    title = sp.get('title') or f"Flange Supplier in {meta['city']} | Remax Forge"
    meta_desc = sp.get('meta_desc') or f"Forged flanges supplied directly from Mumbai to {meta['city']}."
    h1 = sp.get('h1') or f"Flange Supplier in {meta['city']}"
    subline = sp.get('subline') or f"Flanges manufacturer in Mumbai supplying {meta['city']}."
    canonical = sp.get('canonical') or f"https://remaxforge.com{meta['url']}"
    word_floor = sp.get('word_floor') or 1000

    item = {
        'slug': meta['slug'],
        'url': meta['url'],
        'transit': transit_map.get(key, '3-5 days'),
        'price_bands': price_bands_map.get(key, 'ASTM A105 carbon steel: ₹115–140/kg; SS 304L/316L: ₹380–580/kg; Alloy steel F11/F22: ₹230–330/kg'),
        'state': meta['state'],
        'state_hub_url': meta['state_hub_url'],
        'city': meta['city'],
        'floor': word_floor,
        'photo': meta['photo'],
        'video': meta['video'],
        'name': meta['city'],
        'canonical': canonical,
        'title': clean_text(title, key),
        'meta_desc': clean_text(meta_desc, key),
        'h1': clean_text(h1, key),
        'subline': clean_text(subline, key),
        'schema': page_schemas.get(key) or existing_data.get(key, {}).get('schema', {}),
        'sections': page_sections.get(key, []),
        'shared_sections': shared_sections,
        'local_demand_table': local_demand_tables.get(key, []),
        'product_range_table': product_range_tables.get(key, []),
        'faqs': page_faqs.get(key, []),
        'related_links': meta['related_links']
    }
    final_output[key] = item

js_content = "export const cityStrategyData = " + json.dumps(final_output, indent=2, ensure_ascii=False) + ";\n"

with open('frontend/src/assets/data/cityStrategyData.js', 'w', encoding='utf-8') as f:
    f.write(js_content)

print("SUCCESS: frontend/src/assets/data/cityStrategyData.js written.")
for k, v in final_output.items():
    print(f"[{k}] H1: '{v['h1']}' | Subline: '{v['subline'][:40]}...' | Demand rows: {len(v['local_demand_table'])} | Range rows: {len(v['product_range_table'])} | Sections: {len(v['sections'])} | FAQs: {len(v['faqs'])}")
