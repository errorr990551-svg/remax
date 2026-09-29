import openpyxl
import json
import os

wb = openpyxl.load_workbook('frontend/RemaxForge_City_Pages_Content_Strategy_AP_Arunachal.xlsx', data_only=True)

with open('frontend/src/assets/data/cityStrategyData.js', 'r', encoding='utf-8') as f:
    c = f.read().strip()
    if c.startswith('export const cityStrategyData ='):
        c = c[len('export const cityStrategyData ='):].strip()
    if c.endswith(';'):
        c = c[:-1].strip()
    current_data = json.loads(c)

print("=== CHECK 1: PAGES IN DATA VS ON-PAGE SPECS (Tab 6) ===")
s6 = wb['6. On-Page Specs']
tab6_keys = []
for r in range(5, s6.max_row + 1):
    k = s6.cell(r, 1).value
    if k:
        tab6_keys.append(k)
        title = s6.cell(r, 6).value
        meta = s6.cell(r, 9).value
        h1 = s6.cell(r, 12).value
        subline = s6.cell(r, 13).value
        print(f"Tab 6 Key: {k}")
        print(f"  In current_data? {k in current_data}")
        if k in current_data:
            cd = current_data[k]
            print(f"  Title match? {cd.get('title') == title} (Current: '{cd.get('title')[:30]}...' vs Tab6: '{str(title)[:30]}...')")
            print(f"  H1 match? {cd.get('h1') == h1} (Current: '{cd.get('h1')}' vs Tab6: '{h1}')")

print("\n=== CHECK 2: CITY CONTENT SECTIONS (Tab 8) ===")
s8 = wb['8. City Content (Paste)']
tab8_pages = {}
for r in range(5, s8.max_row + 1):
    k = s8.cell(r, 1).value
    sec_num = s8.cell(r, 3).value
    if k and sec_num is not None:
        tab8_pages.setdefault(k, []).append((sec_num, s8.cell(r, 4).value, s8.cell(r, 6).value))

for k, secs in tab8_pages.items():
    if k == 'SHARED':
        print(f"SHARED has {len(secs)} sections in Tab 8: {[s[0] for s in secs]}")
    else:
        in_data = k in current_data
        curr_secs = len(current_data[k].get('sections', [])) if in_data else 0
        print(f"Page '{k}': Tab 8 has {len(secs)} sections, current_data has {curr_secs} sections")
        if in_data:
            c_sec_nums = [float(s['sec_num']) for s in current_data[k].get('sections', [])]
            t_sec_nums = [float(s[0]) if isinstance(s[0], (int, float)) else float(str(s[0])) for s in secs]
            missing = set(t_sec_nums) - set(c_sec_nums)
            if missing:
                print(f"  Missing sections in {k}: {missing}")

print("\n=== CHECK 3: PAGE TABLES (Tab 8b) ===")
s8b = wb['8b. Page Tables (Paste)']
tab8b_counts = {}
for r in range(5, s8b.max_row + 1):
    k = s8b.cell(r, 1).value
    tbl = s8b.cell(r, 2).value
    if k and tbl:
        tab8b_counts.setdefault(k, {}).setdefault(tbl, 0)
        tab8b_counts[k][tbl] += 1

for k, tbls in tab8b_counts.items():
    print(f"Page '{k}': Tab 8b has {tbls}")
    if k in current_data:
        curr_demand = len(current_data[k].get('local_demand_table', []))
        curr_range = len(current_data[k].get('product_range_table', []))
        print(f"  current_data has: demand={curr_demand}, range={curr_range}")

print("\n=== CHECK 4: FAQS (Tab 9) ===")
s9 = wb['9. FAQs (Paste)']
tab9_counts = {}
for r in range(5, s9.max_row + 1):
    k = s9.cell(r, 1).value
    q = s9.cell(r, 4).value
    if k and q:
        tab9_counts[k] = tab9_counts.get(k, 0) + 1

for k, c_count in tab9_counts.items():
    curr_faqs = len(current_data.get(k, {}).get('faqs', []))
    print(f"Page '{k}': Tab 9 has {c_count} FAQs, current_data has {curr_faqs} FAQs")

print("\n=== CHECK 5: LEAD CAPTURE FORMS (Tab 10) ===")
s10 = wb['10. Lead Capture Forms']
print(f"Tab 10 max_row: {s10.max_row}")
for r in range(1, 20):
    vals = [str(s10.cell(r, c).value or '') for c in range(1, 6)]
    if any(vals):
        print(f"  R{r}: {' | '.join(vals)}")

print("\n=== CHECK 6: SCHEMA JSON-LD (Tab 11) ===")
s11 = wb['11. Schema JSON-LD']
for r in range(5, s11.max_row + 1):
    k = s11.cell(r, 1).value
    raw_schema = s11.cell(r, 4).value
    if k:
        print(f"Page '{k}' in Tab 11: Schema present? {bool(raw_schema)} (len {len(str(raw_schema)) if raw_schema else 0})")
        if k in current_data:
            has_s = bool(current_data[k].get('schema'))
            print(f"  In current_data schema? {has_s}")

print("\n=== CHECK 7: MEDIA PLAN (Tab 13) ===")
s13 = wb['13. Media Plan']
print("Photos in Tab 13:")
for r in range(6, 21):
    p_id = s13.cell(r, 1).value
    shot = s13.cell(r, 2).value
    used = s13.cell(r, 8).value
    fname = s13.cell(r, 6).value
    if p_id:
        print(f"  {p_id}: {shot} | file: {fname} | used on: {used}")

print("Videos in Tab 13:")
for r in range(25, 34):
    v_id = s13.cell(r, 1).value
    v_title = s13.cell(r, 2).value
    v_used = s13.cell(r, 7).value
    if v_id:
        print(f"  {v_id}: {v_title} | used on: {v_used}")

print("\n=== CHECK 8: REFERENCE DATA (Tab 14) ===")
s14 = wb['14. Reference Data']
for r in range(1, 15):
    vals = [str(s14.cell(r, c).value or '') for c in range(1, 5)]
    if any(vals):
        print(f"  R{r}: {' | '.join(vals)}")

