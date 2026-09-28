import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('RemaxForge_City_Pages_Content_Strategy_AP_Arunachal.xlsx', data_only=True)
s8 = wb['8. City Content (Paste)']

sections_by_key = {}
for r in range(5, s8.max_row + 1):
    key = s8.cell(r, 1).value
    if not key:
        continue
    sec_num = str(s8.cell(r, 3).value or '')
    sec_name = str(s8.cell(r, 4).value or '')
    elem = str(s8.cell(r, 5).value or '')
    heading = str(s8.cell(r, 6).value or '')
    copy = str(s8.cell(r, 7).value or '')
    words = s8.cell(r, 8).value
    
    if key not in sections_by_key:
        sections_by_key[key] = []
    sections_by_key[key].append({
        'sec_num': sec_num,
        'sec_name': sec_name,
        'element': elem,
        'heading': heading,
        'copy': copy,
        'words': words
    })

for k, secs in sections_by_key.items():
    nums = [s['sec_num'] + ' ' + s['sec_name'] for s in secs]
    print(f'Key {k}: {len(secs)} sections -> ' + '; '.join(nums))

print('\n*** Checking 8b. Page Tables (Paste) ***')
s8b = wb['8b. Page Tables (Paste)']
tables_by_key = {}
for r in range(5, s8b.max_row + 1):
    key = s8b.cell(r, 1).value
    if not key:
        continue
    tbl = str(s8b.cell(r, 2).value or '')
    row_num = s8b.cell(r, 3).value
    c1 = s8b.cell(r, 4).value
    c2 = s8b.cell(r, 5).value
    c3 = s8b.cell(r, 6).value
    c4 = s8b.cell(r, 7).value
    if key not in tables_by_key:
        tables_by_key[key] = {}
    if tbl not in tables_by_key[key]:
        tables_by_key[key][tbl] = []
    tables_by_key[key][tbl].append([c1, c2, c3, c4])

for k, tbls in tables_by_key.items():
    print(f'Key {k}: tables -> ' + ', '.join([f'{t} ({len(rows)} rows)' for t, rows in tbls.items()]))

print('\n*** Checking 9. FAQs (Paste) ***')
s9 = wb['9. FAQs (Paste)']
faqs_by_key = {}
for r in range(5, s9.max_row + 1):
    key = s9.cell(r, 1).value
    if not key:
        continue
    q = s9.cell(r, 4).value
    a = s9.cell(r, 5).value
    if key not in faqs_by_key:
        faqs_by_key[key] = []
    faqs_by_key[key].append({'question': q, 'answer': a})

for k, faqs in faqs_by_key.items():
    print(f'Key {k}: {len(faqs)} FAQs')
