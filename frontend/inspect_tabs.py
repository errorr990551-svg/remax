import openpyxl

wb = openpyxl.load_workbook('frontend/RemaxForge_City_Pages_Content_Strategy_AP_Arunachal.xlsx', data_only=True)

with open('frontend/tabs_audit.txt', 'w', encoding='utf-8') as out:
    for name in wb.sheetnames:
        ws = wb[name]
        out.write(f"=== SHEET: {name} ===\n")
        out.write(f"Dimensions: {ws.max_row} rows, {ws.max_column} cols\n")
        for r in range(1, min(7, ws.max_row + 1)):
            vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
            if any(vals):
                out.write(f"R{r}: " + " | ".join(vals[:12]) + "\n")
        out.write("\n")

print("Done! Check frontend/tabs_audit.txt")
